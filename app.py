import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io

st.set_page_config(page_title="QA Test Matrix Generator", page_icon="🧪", layout="centered")

st.title("🧪 QA Test Matrix Generator")
st.write("Upload your QA Test Matrix Excel files (`.xlsx` / `.xls`) to process, format, or generate standardized QA Test Case matrices.")

# Excel File Uploader
uploaded_files = st.file_uploader(
    "Upload QA Matrix Excel Files (.xlsx / .xls)",
    type=["xlsx", "xls"],
    accept_multiple_files=True
)

st.subheader("Matrix Metadata Options")
col1, col2, col3 = st.columns(3)

with col1:
    # Application Selector: MVA or MVO
    application_type = st.selectbox(
        "Application",
        options=["MVA", "MVO"],
        index=0,
        help="MVA = My Verizon App | MVO = My Verizon Online"
    )

with col2:
    # Label is 'Test Case Owner Name' -> Populates the 'Assigned To' column in output
    tc_owner_name = st.text_input("Test Case Owner Name", value="Sawan")

with col3:
    environment = st.text_input("Environment", value="QA1")

if uploaded_files:
    if st.button("🚀 Process & Generate Test Matrix", type="primary"):
        all_dfs = []
        
        for file in uploaded_files:
            try:
                # Read uploaded excel sheet
                df = pd.read_excel(file)
                all_dfs.append(df)
            except Exception as e:
                st.error(f"Error reading {file.name}: {e}")

        if all_dfs:
            combined_df = pd.concat(all_dfs, ignore_index=True)
            
            # 1. Enforce STATUS = "New" for ALL rows
            if "STATUS" in combined_df.columns:
                combined_df["STATUS"] = "New"
            elif "Status" in combined_df.columns:
                combined_df["Status"] = "New"
            else:
                combined_df["STATUS"] = "New"

            # 2. Update APPLICATION column (MVA / MVO)
            if "APPLICATION" in combined_df.columns:
                combined_df["APPLICATION"] = application_type
            elif "Application" in combined_df.columns:
                combined_df["Application"] = application_type

            # 3. Map 'Test Case Owner Name' directly to 'ASSIGNED TO' / 'Assigned To' column
            if "ASSIGNED TO" in combined_df.columns:
                combined_df["ASSIGNED TO"] = tc_owner_name
            elif "Assigned To" in combined_df.columns:
                combined_df["Assigned To"] = tc_owner_name
            elif "Assigned to" in combined_df.columns:
                combined_df["Assigned to"] = tc_owner_name

            # 4. Update ENVIRONMENT column
            if "ENVIRONMENT" in combined_df.columns:
                combined_df["ENVIRONMENT"] = environment
            elif "Environment" in combined_df.columns:
                combined_df["Environment"] = environment

            # Generate formatted openpyxl Excel file in memory
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                combined_df.to_excel(writer, index=False, sheet_name="QA_Test_Matrix")
            
            output.seek(0)
            
            # Apply formatting and styling
            wb = openpyxl.load_workbook(output)
            ws = wb.active
            
            header_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True)
            thin_border = Border(
                left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC')
            )
            
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = thin_border
                
            for r_idx in range(2, ws.max_row + 1):
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=r_idx, column=col_idx)
                    cell.font = Font(name="Calibri", size=11)
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    cell.border = thin_border
                    
            for col in ws.columns:
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = 25
                
            styled_output = io.BytesIO()
            wb.save(styled_output)
            styled_output.seek(0)

            st.success(f"Successfully processed {len(combined_df)} rows for **{application_type}** assigned to **{tc_owner_name}**!")
            
            st.download_button(
                label=f"📥 Download {application_type} Excel Matrix (.xlsx)",
                data=styled_output,
                file_name=f"{application_type}_QA_Test_Matrix.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
